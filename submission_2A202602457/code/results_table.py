"""results_table.py — lưu kết quả từng lần chạy ra JSON, rồi điền vào experiments.xlsx từ mẫu
templates/experiment_table_template.xlsx (giữ nguyên 4 sheet, tên cột và các cột công thức).
"""
from __future__ import annotations

import json
import math
from pathlib import Path

COLUMNS = ["exp_id", "group", "description", "loss", "optimizer", "lr", "weight_decay", "batch", "epochs",
           "hidden", "dropout", "clip_norm", "precision", "init", "seed", "step0_loss", "best_val_loss",
           "best_epoch", "final_train_loss", "final_val_loss", "val_acc", "val_macro_f1", "time_per_epoch_s",
           "peak_mem_MB", "diverged", "eval_acc", "eval_macro_f1", "figure_file", "notes"]
FORMULA_COLUMNS = {"step0_gap_vs_lnC", "gap_val_minus_train", "delta_val_f1_vs_base", "beyond_noise"}

LOSS_NAMES = {"ce": "CE", "mse": "MSE"}
OPT_NAMES = {"sgd": "SGD", "sgd_momentum": "SGD+momentum", "adam": "Adam", "adamw": "AdamW"}


def _clean(x):
    """Đổi NaN/inf thành None để JSON hợp lệ và ô Excel để trống."""
    if isinstance(x, float) and not math.isfinite(x):
        return None
    if isinstance(x, (list, tuple)):
        return [_clean(v) for v in x]
    if isinstance(x, dict):
        return {k: _clean(v) for k, v in x.items()}
    return x


def save_result(result: dict, results_dir: str = "../results") -> str:
    """Ghi cfg, history, summary (KHÔNG ghi best_state) ra <results_dir>/<exp_id>.json."""
    Path(results_dir).mkdir(parents=True, exist_ok=True)
    path = Path(results_dir) / f"{result['cfg']['exp_id']}.json"
    payload = {k: _clean(result[k]) for k in ("cfg", "history", "summary")}
    for k in ("eval", "notes"):
        if k in result:
            payload[k] = _clean(result[k])
    with open(path, "w", encoding="utf-8") as f:
        json.dump(payload, f, indent=1, ensure_ascii=False)
    return str(path)


def load_result(exp_id: str, results_dir: str = "../results") -> dict | None:
    path = Path(results_dir) / f"{exp_id}.json"
    if not path.is_file():
        return None
    with open(path, encoding="utf-8") as f:
        r = json.load(f)
    r["cfg"]["hidden"] = tuple(r["cfg"]["hidden"])
    return r


def load_results(results_dir: str = "../results") -> list[dict]:
    """Đọc mọi file *.json trong results_dir, sắp theo exp_id."""
    return [load_result(p.stem, results_dir) for p in sorted(Path(results_dir).glob("*.json"))]


def to_row(result: dict, eval_scores: dict | None = None, notes: str = "") -> dict:
    """Một kết quả -> một dòng của bảng. Chỉ truyền eval_scores cho baseline và cấu hình cuối cùng."""
    cfg, s = result["cfg"], result["summary"]
    note_parts = [n for n in (result.get("notes", ""), notes) if n]
    if cfg.get("scheduler"):
        note_parts.append(f"scheduler={cfg['scheduler']}")
    if cfg["group"] in ("init", "baseline") and s.get("act_std_step0"):
        note_parts.append("std kích hoạt bước 0 (ReLU1, ReLU2, logit) = "
                          + ", ".join(f"{v:.4g}" for v in s["act_std_step0"]))
    if cfg["optimizer"] in ("sgd_momentum",):
        note_parts.append(f"momentum={cfg['momentum']}")
    if cfg["optimizer"] in ("adam", "adamw"):
        note_parts.append("betas=(0.9,0.999), eps=1e-8")
    row = dict(
        exp_id=cfg["exp_id"], group=cfg["group"], description=cfg["description"],
        loss=LOSS_NAMES[cfg["loss"]], optimizer=OPT_NAMES[cfg["optimizer"]], lr=cfg["lr"],
        weight_decay=cfg["weight_decay"], batch=cfg["batch"], epochs=cfg["epochs"],
        hidden="-".join(str(h) for h in cfg["hidden"]), dropout=cfg["dropout"],
        clip_norm=cfg["clip_norm"] if cfg["clip_norm"] is not None else "none",
        precision=cfg["precision"], init=cfg["init"], seed=cfg["seed"],
        step0_loss=s["step0_loss"], best_val_loss=s["best_val_loss"], best_epoch=s["best_epoch"],
        final_train_loss=s["final_train_loss"], final_val_loss=s["final_val_loss"],
        val_acc=s["val_acc"], val_macro_f1=s["val_macro_f1"], time_per_epoch_s=s["time_per_epoch_s"],
        peak_mem_MB=s["peak_mem_MB"], diverged="yes" if s["diverged"] else "no",
        eval_acc=None, eval_macro_f1=None,
        figure_file=f"figures/{cfg['exp_id']}.png", notes="; ".join(note_parts),
    )
    if eval_scores is not None:
        row["eval_acc"] = eval_scores["accuracy"]
        row["eval_macro_f1"] = eval_scores["macro_f1"]
    return {k: _clean(v) for k, v in row.items()}


def write_xlsx(rows: list[dict], template_path: str, out_path: str,
               seed_ids: list[str] | None = None, summary_notes: dict | None = None) -> None:
    """Điền các dòng vào sheet "Experiments" từ dòng 2, các exp_id seed vào sheet "Seeds",
    nhận xét vào cột H của sheet "Summary"; giữ nguyên mọi ô công thức."""
    import openpyxl

    wb = openpyxl.load_workbook(template_path)  # KHÔNG data_only=True (sẽ mất công thức)
    ws = wb["Experiments"]
    header = {c.value: c.column for c in ws[1] if c.value}
    assert all(c in header for c in COLUMNS), "mẫu thiếu cột"
    max_rows = ws.max_row - 1
    assert len(rows) <= max_rows, f"mẫu chỉ có công thức cho {max_rows} dòng, nhận {len(rows)}"
    # xoá giá trị điền sẵn (không phải công thức) của các dòng dữ liệu
    for r in range(2, ws.max_row + 1):
        for name in COLUMNS:
            ws.cell(row=r, column=header[name]).value = None
    for i, row in enumerate(rows):
        for name in COLUMNS:
            ws.cell(row=2 + i, column=header[name]).value = row.get(name)

    if seed_ids is not None:
        seeds = wb["Seeds"]
        assert len(seed_ids) <= 5, "sheet Seeds có chỗ cho tối đa 5 seed"
        for i in range(5):
            seeds.cell(row=2 + i, column=1).value = seed_ids[i] if i < len(seed_ids) else None

    if summary_notes:
        summ = wb["Summary"]
        for r in range(2, summ.max_row + 1):
            g = summ.cell(row=r, column=1).value
            if g in summary_notes:
                summ.cell(row=r, column=8).value = summary_notes[g]

    # openpyxl không lưu giá trị đã tính của công thức: buộc Excel/LibreOffice tính lại khi mở
    wb.calculation.fullCalcOnLoad = True
    Path(out_path).parent.mkdir(parents=True, exist_ok=True)
    wb.save(out_path)
